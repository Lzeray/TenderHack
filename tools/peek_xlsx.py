"""Показать первые строки выгрузок СТЕ (Excel) — только для изучения формулировок."""
import sys
import openpyxl

for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    print("==", path, ws.max_row, ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        print([str(x)[:300] if x is not None else None for x in row])
        if i >= int(__import__("os").environ.get("PEEK_ROWS", "6")):
            break
